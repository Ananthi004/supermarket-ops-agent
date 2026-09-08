from pathlib import Path
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


# Folder where generated invoices are stored
INVOICE_DIR = Path("data/invoices")
INVOICE_DIR.mkdir(parents=True, exist_ok=True)


def generate_invoice_pdf(bill_id: int) -> dict:
    """
    Generate a GST invoice PDF for a finalized bill.
    """

    from tools.billing import get_bill

    bill_result = get_bill(bill_id)

    if not bill_result.get("success"):
        return {
            "success": False,
            "message": bill_result.get("message", "Bill not found.")
        }

    bill = bill_result["bill"]

    if bill["status"] != "finalized":
        return {
            "success": False,
            "message": "Invoice can only be generated for a finalized bill."
        }

    items = bill.get("items", [])

    if not items:
        return {
            "success": False,
            "message": "Cannot generate invoice because the bill has no items."
        }

    file_path = INVOICE_DIR / f"invoice_bill_{bill_id}.pdf"

    try:
        document = SimpleDocTemplate(
            str(file_path),
            pagesize=A4,
            rightMargin=35,
            leftMargin=35,
            topMargin=35,
            bottomMargin=35,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "InvoiceTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            fontSize=18,
            spaceAfter=8,
        )

        center_style = ParagraphStyle(
            "Center",
            parent=styles["Normal"],
            alignment=TA_CENTER,
        )

        right_style = ParagraphStyle(
            "Right",
            parent=styles["Normal"],
            alignment=TA_RIGHT,
        )

        story = []

        # -------------------------------------------------
        # HEADER
        # -------------------------------------------------

        story.append(Paragraph("SUPERMARKET", title_style))
        story.append(
            Paragraph(
                "GST TAX INVOICE",
                ParagraphStyle(
                    "GSTTitle",
                    parent=styles["Heading2"],
                    alignment=TA_CENTER,
                ),
            )
        )

        story.append(Spacer(1, 12))

        created_at = bill.get("created_at", "")
        if created_at:
            try:
                date_text = datetime.fromisoformat(
                    created_at.replace("Z", "")
                ).strftime("%d-%m-%Y %I:%M %p")
            except Exception:
                date_text = str(created_at)
        else:
            date_text = datetime.now().strftime("%d-%m-%Y %I:%M %p")

        customer_name = bill.get("customer_name") or "Walk-in Customer"

        bill_info = [
            [
                Paragraph("<b>Bill No.</b>", styles["Normal"]),
                str(bill_id),
                Paragraph("<b>Date</b>", styles["Normal"]),
                date_text,
            ],
            [
                Paragraph("<b>Customer</b>", styles["Normal"]),
                customer_name,
                Paragraph("<b>Payment</b>", styles["Normal"]),
                str(bill.get("payment_mode") or "").upper(),
            ],
        ]

        info_table = Table(
            bill_info,
            colWidths=[70, 170, 70, 180],
        )

        info_table.setStyle(
            TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ])
        )

        story.append(info_table)
        story.append(Spacer(1, 15))

        # -------------------------------------------------
        # ITEMS
        # -------------------------------------------------

        item_data = [
            [
                "S.No.",
                "Product",
                "Qty",
                "Unit Price",
                "GST %",
                "GST",
                "Amount",
            ]
        ]

        for index, item in enumerate(items, start=1):
            item_data.append([
                str(index),
                item["product_name"],
                f'{item["quantity"]:.2f}',
                f'₹{item["unit_price"]:.2f}',
                f'{item["gst_rate"]:.2f}%',
                f'₹{item["gst_amount"]:.2f}',
                f'₹{item["line_total"]:.2f}',
            ])

        item_table = Table(
            item_data,
            colWidths=[35, 145, 45, 75, 55, 60, 70],
            repeatRows=1,
        )

        item_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ALIGN", (0, 0), (0, -1), "CENTER"),
                ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ])
        )

        story.append(item_table)
        story.append(Spacer(1, 15))

        # -------------------------------------------------
        # GST SUMMARY
        # -------------------------------------------------

        summary_data = [
            ["Subtotal", f'₹{bill["subtotal"]:.2f}'],
            ["CGST", f'₹{bill["cgst_amount"]:.2f}'],
            ["SGST", f'₹{bill["sgst_amount"]:.2f}'],
            ["Total GST", f'₹{bill["gst_amount"]:.2f}'],
            ["Grand Total", f'₹{bill["total_amount"]:.2f}'],
        ]

        summary_table = Table(
            summary_data,
            colWidths=[380, 110],
            hAlign="RIGHT",
        )

        summary_table.setStyle(
            TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ])
        )

        story.append(summary_table)
        story.append(Spacer(1, 15))

        # -------------------------------------------------
        # PAYMENT DETAILS
        # -------------------------------------------------

        payment_amount = bill.get("payment_amount", 0) or 0

        payment_data = [
            [
                Paragraph("<b>Payment Mode</b>", styles["Normal"]),
                str(bill.get("payment_mode") or "").upper(),
            ],
            [
                Paragraph("<b>Amount Paid</b>", styles["Normal"]),
                f"₹{payment_amount:.2f}",
            ],
        ]

        if bill.get("payment_reference"):
            payment_data.append([
                Paragraph("<b>Reference</b>", styles["Normal"]),
                str(bill["payment_reference"]),
            ])

        payment_table = Table(
            payment_data,
            colWidths=[130, 360],
        )

        payment_table.setStyle(
            TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ])
        )

        story.append(payment_table)
        story.append(Spacer(1, 25))

        story.append(
            Paragraph(
                "Thank you for shopping with us!",
                center_style
            )
        )

        story.append(Spacer(1, 8))

        story.append(
            Paragraph(
                "This is a computer-generated invoice.",
                center_style
            )
        )

        # -------------------------------------------------
        # BUILD PDF
        # -------------------------------------------------

        document.build(story)

        return {
            "success": True,
            "message": f"Invoice generated successfully for Bill #{bill_id}.",
            "bill_id": bill_id,
            "file_path": str(file_path),
        }

    except Exception as e:
        return {
            "success": False,
            "message": f"Failed to generate invoice: {str(e)}"
        }

from pptx import Presentation
from pptx.util import Inches, Pt


def generate_daily_analysis_pptx() -> dict:
    """
    Generate a PowerPoint analysis deck for today's supermarket sales.
    """

    from tools.analytics import get_daily_sales

    analytics = get_daily_sales()

    if not analytics.get("success"):
        return {
            "success": False,
            "message": analytics.get(
                "message",
                "Unable to generate analytics."
            )
        }

    try:
        date = analytics["date"]

        file_path = (
            Path("data/analysis")
            / f"daily_analysis_{date}.pptx"
        )

        file_path.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        prs = Presentation()

        # -------------------------------------------------
        # SLIDE 1 - TITLE
        # -------------------------------------------------

        slide = prs.slides.add_slide(
            prs.slide_layouts[0]
        )

        slide.shapes.title.text = "Supermarket Daily Analysis"

        slide.placeholders[1].text = (
            f"Sales Report - {date}"
        )

        # -------------------------------------------------
        # SLIDE 2 - SALES SUMMARY
        # -------------------------------------------------

        slide = prs.slides.add_slide(
            prs.slide_layouts[5]
        )

        slide.shapes.title.text = "Daily Sales Summary"

        summary = [
            ["Metric", "Value"],
            ["Total Bills", str(analytics["bill_count"])],
            ["Subtotal", f'₹{analytics["subtotal"]:.2f}'],
            ["GST", f'₹{analytics["gst_amount"]:.2f}'],
            ["CGST", f'₹{analytics["cgst_amount"]:.2f}'],
            ["SGST", f'₹{analytics["sgst_amount"]:.2f}'],
            ["Total Sales", f'₹{analytics["total_sales"]:.2f}'],
        ]

        table = slide.shapes.add_table(
            len(summary),
            2,
            Inches(1),
            Inches(1.7),
            Inches(8),
            Inches(4)
        ).table

        for r, row in enumerate(summary):
            for c, value in enumerate(row):
                table.cell(r, c).text = value

        # -------------------------------------------------
        # SLIDE 3 - PAYMENT ANALYSIS
        # -------------------------------------------------

        slide = prs.slides.add_slide(
            prs.slide_layouts[5]
        )

        slide.shapes.title.text = "Payment Mode Analysis"

        payments = analytics["payments"]

        payment_data = [
            ["Payment Mode", "Sales", "Collected"],
            [
                "Cash",
                f'₹{payments["cash"]["sales"]:.2f}',
                f'₹{payments["cash"]["collected"]:.2f}'
            ],
            [
                "UPI",
                f'₹{payments["upi"]["sales"]:.2f}',
                f'₹{payments["upi"]["collected"]:.2f}'
            ],
            [
                "Card",
                f'₹{payments["card"]["sales"]:.2f}',
                f'₹{payments["card"]["collected"]:.2f}'
            ],
            [
                "Credit",
                f'₹{payments["credit"]["sales"]:.2f}',
                f'₹{payments["credit"]["collected"]:.2f}'
            ],
        ]

        table = slide.shapes.add_table(
            len(payment_data),
            3,
            Inches(0.7),
            Inches(1.7),
            Inches(8.8),
            Inches(4)
        ).table

        for r, row in enumerate(payment_data):
            for c, value in enumerate(row):
                table.cell(r, c).text = value

        # -------------------------------------------------
        # SLIDE 4 - TOP PRODUCTS
        # -------------------------------------------------

        slide = prs.slides.add_slide(
            prs.slide_layouts[5]
        )

        slide.shapes.title.text = "Top Selling Products"

        top_products = analytics["top_products"]

        if top_products:

            product_data = [
                ["Product", "Quantity Sold", "Sales"]
            ]

            for product in top_products:
                product_data.append([
                    product["product_name"],
                    str(product["quantity_sold"]),
                    f'₹{product["sales_amount"]:.2f}'
                ])

            table = slide.shapes.add_table(
                len(product_data),
                3,
                Inches(0.7),
                Inches(1.5),
                Inches(8.8),
                Inches(4.5)
            ).table

            for r, row in enumerate(product_data):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value

        else:

            textbox = slide.shapes.add_textbox(
                Inches(1),
                Inches(2),
                Inches(8),
                Inches(1)
            )

            textbox.text_frame.text = (
                "No products were sold today."
            )

        # -------------------------------------------------
        # SLIDE 5 - LOW STOCK
        # -------------------------------------------------

        slide = prs.slides.add_slide(
            prs.slide_layouts[5]
        )

        slide.shapes.title.text = "Low Stock Alert"

        low_stock = analytics["low_stock_products"]

        if low_stock:

            stock_data = [
                [
                    "Product",
                    "Current Stock",
                    "Reorder Level"
                ]
            ]

            for product in low_stock:

                stock_data.append([
                    product["product_name"],
                    f'{product["quantity"]:.2f} '
                    f'{product["unit"]}',
                    f'{product["reorder_level"]:.2f}'
                ])

            table = slide.shapes.add_table(
                len(stock_data),
                3,
                Inches(0.7),
                Inches(1.5),
                Inches(8.8),
                Inches(4.5)
            ).table

            for r, row in enumerate(stock_data):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value

        else:

            textbox = slide.shapes.add_textbox(
                Inches(1),
                Inches(2),
                Inches(8),
                Inches(1)
            )

            textbox.text_frame.text = (
                "No low-stock products."
            )

        # -------------------------------------------------
        # SLIDE 6 - OWNER RECOMMENDATIONS
        # -------------------------------------------------

        slide = prs.slides.add_slide(
            prs.slide_layouts[5]
        )

        slide.shapes.title.text = "Owner Recommendations"

        recommendations = []

        if low_stock:
            recommendations.append(
                "Restock products that are at or below reorder level."
            )

        if top_products:
            recommendations.append(
                f"Focus on fast-moving product: "
                f"{top_products[0]['product_name']}."
            )

        if payments["credit"]["sales"] > 0:
            recommendations.append(
                "Monitor customer credit and Khata outstanding balances."
            )

        if not recommendations:
            recommendations.append(
                "Continue monitoring daily sales and inventory."
            )

        textbox = slide.shapes.add_textbox(
            Inches(1),
            Inches(1.7),
            Inches(8),
            Inches(4.5)
        )

        text_frame = textbox.text_frame

        for index, recommendation in enumerate(recommendations):

            if index == 0:
                paragraph = text_frame.paragraphs[0]
            else:
                paragraph = text_frame.add_paragraph()

            paragraph.text = f"• {recommendation}"
            paragraph.font.size = Pt(20)

        # -------------------------------------------------
        # SAVE
        # -------------------------------------------------

        prs.save(str(file_path))

        return {
            "success": True,
            "message": (
                f"Daily analysis PPTX generated successfully "
                f"for {date}."
            ),
            "file_path": str(file_path),
            "date": date,
            "slides": 6
        }

    except Exception as e:

        return {
            "success": False,
            "message": f"Failed to generate PPTX: {str(e)}"
        }